# -*- coding: utf-8 -*-
"""
Builds an Executive-Grade Microsoft Word Document (.docx) for:
VEO Rule-Based Error Classification: Comprehensive Technical Reference & Explanation Manual
Target: docs/VEO_RULE_BASED_CLASSIFICATION_EXPLAINED.docx
"""

import os
import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding="utf-8")

MD_PATH = Path("docs/VEO_RULE_BASED_CLASSIFICATION_EXPLAINED.md")
DOCX_PATH = Path("docs/VEO_RULE_BASED_CLASSIFICATION_EXPLAINED.docx")

# ==============================================================================
# OpenXML Styling Utilities
# ==============================================================================

def set_cell_background(cell, fill_hex):
    """Sets hex background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Sets cell padding in dxa (1 pt = 20 dxa). Default 4pt top/bottom, 6pt left/right."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, border_color="CBD5E1", header_color="0F2942"):
    """Applies clean executive table borders (thick top/bottom rules, thin horizontal inside, no vertical)."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '10')  # 1.25 pt
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), header_color)
    tblBorders.append(top)
    
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), header_color)
    tblBorders.append(bottom)
    
    insideH = OxmlElement('w:insideH')
    insideH.set(qn('w:val'), 'single')
    insideH.set(qn('w:sz'), '4')  # 0.5 pt
    insideH.set(qn('w:space'), '0')
    insideH.set(qn('w:color'), border_color)
    tblBorders.append(insideH)
    
    for b_name in ['left', 'right', 'insideV']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tblBorders.append(b)
        
    tblPr.append(tblBorders)

def set_callout_borders(table, border_color="2563EB"):
    """Thick left border for callout boxes; no top/right/bottom/inside borders."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3.0 pt thick accent bar
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tblBorders.append(left)
    
    for b_name in ['top', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tblBorders.append(b)
        
    tblPr.append(tblBorders)

def set_code_borders(table, border_color="CBD5E1"):
    """Thin border around code block."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b_name in ['top', 'bottom', 'left', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')  # 0.5 pt
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), border_color)
        tblBorders.append(b)
    for b_name in ['insideH', 'insideV']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tblBorders.append(b)
    tblPr.append(tblBorders)

def prevent_row_split(row):
    """Prevents row from splitting across a page boundary."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def set_repeat_header(row):
    """Ensures table header repeats at top of next page if table spans pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

# ==============================================================================
# Text Formatting Helpers
# ==============================================================================

def clean_latex(text: str) -> str:
    """Converts common LaTeX notations into clean unicode text for Word rendering."""
    t = text
    # Replace math wrappers
    t = t.replace(r"\text{", "").replace(r"\mathrm{", "")
    t = t.replace(r"\mathbb{R}", "ℝ")
    t = t.replace(r"\times", "×")
    t = t.replace(r"\cdot", "·")
    t = t.replace(r"\ge", "≥").replace(r"\le", "≤")
    t = t.replace(r"\neq", "≠").replace(r"\approx", "≈")
    t = t.replace(r"\in", "∈").replace(r"\subset", "⊂")
    t = t.replace(r"\cup", "∪").replace(r"\cap", "∩")
    t = t.replace(r"\land", " ∧ ").replace(r"\lor", " ∨ ")
    t = t.replace(r"\neg", "¬")
    t = t.replace(r"\rightarrow", "→").replace(r"\implies", " ⟹ ")
    t = t.replace(r"\nabla", "∇")
    t = t.replace(r"\kappa", "κ")
    t = t.replace(r"\tau", "τ")
    t = t.replace(r"\mu", "μ")
    t = t.replace(r"\sigma", "σ")
    t = t.replace(r"\Delta", "Δ")
    t = t.replace(r"\mathcal{L}", "ℒ")
    t = t.replace(r"\mathcal{R}", "ℛ")
    t = t.replace(r"\infty", "∞")
    t = t.replace(r"\_", "_")
    # Clean up braces and math symbols
    t = re.sub(r"\\frac\{([^}]+)\}\{([^}]+)\}", r"(\1 / \2)", t)
    t = re.sub(r"\\sqrt\{([^}]+)\}", r"√(\1)", t)
    t = re.sub(r"\\mathbf\{([^}]+)\}", r"\1", t)
    t = re.sub(r"\\mid", "|", t)
    t = re.sub(r"\\%", "%", t)
    t = t.replace("{", "").replace("}", "")
    return t

def add_formatted_runs(paragraph, text, base_font_size=10, is_bold=False, is_italic=False, color_rgb=(0x1E, 0x29, 0x3B)):
    """Parses markdown inline tokens (**bold**, *italic*, `code`, $math$) and adds formatted runs."""
    # Pattern to match bold, italic, inline code, and inline math
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$[^$]+\$)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run_text = token[2:-2]
            r = paragraph.add_run(clean_latex(run_text))
            r.font.bold = True
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = RGBColor(*color_rgb)
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run_text = token[1:-1]
            r = paragraph.add_run(clean_latex(run_text))
            r.font.italic = True
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = RGBColor(*color_rgb)
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run_text = token[1:-1]
            r = paragraph.add_run(run_text)
            r.font.name = 'Consolas'
            r.font.size = Pt(base_font_size * 0.9)
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)  # Slate 900
        elif token.startswith("$") and token.endswith("$") and len(token) >= 2:
            run_text = token[1:-1]
            clean_math = clean_latex(run_text)
            r = paragraph.add_run(clean_math)
            r.font.name = 'Cambria Math'
            r.font.italic = True
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # Deep Indigo
        else:
            r = paragraph.add_run(clean_latex(token))
            r.font.bold = is_bold
            r.font.italic = is_italic
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = RGBColor(*color_rgb)

# ==============================================================================
# Document Generation Pipeline
# ==============================================================================

def generate_docx():
    print(f"Reading markdown from {MD_PATH}...")
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()
    
    # Page Setup (Standard Letter, 0.75" Margins -> 7.0" Usable Width)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # Configure Default Styles
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(10)
    normal_font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    # Configure Header & Footer
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header_p.add_run("VISION ERROR OBSERVATORY (VEO) — RULE-BASED CLASSIFICATION MANUAL")
    hrun.font.name = 'Calibri'
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer_p.add_run("NMIMS MPSTME / Hyderabad • Capstone Project 2026–2027 • Confirmed Source-Code Reference")
    frun.font.name = 'Calibri'
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # --------------------------------------------------------------------------
    # COVER / HEADER BLOCK
    # --------------------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("NMIMS MPSTME / NMIMS HYDERABAD • DEPARTMENT OF COMPUTER ENGINEERING")
    r_inst.font.size = Pt(8.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("VEO RULE-BASED ERROR CLASSIFICATION")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)  # Deep Navy
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Comprehensive Technical Reference & Plain-Language Explanation Manual")
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    p_sub.paragraph_format.space_after = Pt(12)
    
    # Metadata Overview Table (2-column executive table)
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_meta = [Inches(3.5), Inches(3.5)]
    set_table_borders(meta_table, border_color="E2E8F0", header_color="CBD5E1")
    
    meta_data = [
        [("Document Version:", " 1.0 (Production / Viva Ready)"), ("Target Audience:", " Beginners, CV Researchers, Examiners")],
        [("Primary Core Modules:", " src/veo_nyu/scoring.py, features.py"), ("Evaluated Dataset:", " NYU Depth V2 (5,236 test regions)")],
        [("Interactive Web Tool:", " web_app/index.html (Port 8000)"), ("Verification Status:", " 100% Code-Audited & Verified")]
    ]
    
    for r_idx, row in enumerate(meta_table.rows):
        prevent_row_split(row)
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_w_meta[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            set_cell_background(cell, "F8FAFC")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            lbl, val = meta_data[r_idx][c_idx]
            r1 = p.add_run(lbl)
            r1.font.bold = True
            r1.font.size = Pt(9)
            r1.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
            r2 = p.add_run(val)
            r2.font.size = Pt(9)
            r2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --------------------------------------------------------------------------
    # PARSING MARKDOWN SECTIONS
    # --------------------------------------------------------------------------
    lines = md_text.splitlines()
    i = 0
    total_lines = len(lines)
    
    # Helper to add section headings
    def add_chapter_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(5)
        h.paragraph_format.keep_with_next = True
        
        clean_text = re.sub(r'^#+\s*', '', text).strip()
        match = re.match(r'^(Chapter \d+:)(.*)$', clean_text)
        if match:
            prefix, rest = match.groups()
            r_pre = h.add_run(prefix + " ")
            r_pre.font.bold = True
            r_pre.font.size = Pt(14)
            r_pre.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)  # Primary Accent Blue
            r_t = h.add_run(rest.strip().upper())
            r_t.font.bold = True
            r_t.font.size = Pt(14)
            r_t.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)  # Deep Navy
        else:
            r = h.add_run(text.replace("#", "").strip())
            r.font.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)

    def add_subheading(text, level=2):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(11 if level == 2 else 8)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        clean_text = re.sub(r'^#+\s*', '', text).strip()
        r = h.add_run(clean_text)
        r.font.bold = True
        r.font.size = Pt(12 if level == 2 else 10.5)
        sub_color = (0x1E, 0x3A, 0x8A) if level == 2 else (0x0F, 0x17, 0x2A)
        r.font.color.rgb = RGBColor(*sub_color)

    while i < total_lines:
        line = lines[i].strip()
        
        # Skip top title and metadata since we already generated executive cover block
        if line.startswith("# VEO Rule-Based Error Classification") or line.startswith("> **Document Version") or line.startswith("> **Status") or line.startswith("> **Target Audience") or line.startswith("> **Primary Source Code") or line.startswith("> **Interactive UI"):
            i += 1
            continue
        
        # Skip empty lines
        if not line:
            i += 1
            continue
            
        # Horizontal rule
        if line == "---":
            # Just add space
            i += 1
            continue
            
        # Headings
        if line.startswith("# Chapter "):
            add_chapter_heading(line)
            i += 1
            continue
        elif line.startswith("## Table of Contents"):
            add_subheading("Table of Contents", level=2)
            i += 1
            continue
        elif line.startswith("### "):
            add_subheading(line, level=3)
            i += 1
            continue
        elif line.startswith("## "):
            add_subheading(line, level=2)
            i += 1
            continue
        elif line.startswith("#### "):
            add_subheading(line, level=3)
            i += 1
            continue
            
        # Callout Alert Boxes (> [!IMPORTANT], > [!WARNING], > [!CAUTION], > [!NOTE])
        if line.startswith(">"):
            alert_lines = []
            while i < total_lines and lines[i].strip().startswith(">"):
                alert_lines.append(lines[i].strip()[1:].strip())
                i += 1
            
            full_alert = "\n".join(alert_lines)
            border_color = "2563EB"
            bg_color = "EFF6FF"
            icon_title = "📌 NOTE"
            
            if "[!IMPORTANT]" in full_alert:
                border_color = "7C3AED"
                bg_color = "F5F3FF"
                icon_title = "⚡ IMPORTANT REQUIREMENT"
                full_alert = full_alert.replace("[!IMPORTANT]", "")
            elif "[!WARNING]" in full_alert:
                border_color = "D97706"
                bg_color = "FFFBEB"
                icon_title = "⚠️ CRITICAL WARNING"
                full_alert = full_alert.replace("[!WARNING]", "")
            elif "[!CAUTION]" in full_alert:
                border_color = "DC2626"
                bg_color = "FEF2F2"
                icon_title = "🚨 CAUTION / BOUNDARY"
                full_alert = full_alert.replace("[!CAUTION]", "")
            elif "[!TIP]" in full_alert:
                border_color = "059669"
                bg_color = "ECFDF5"
                icon_title = "💡 BEST PRACTICE TIP"
                full_alert = full_alert.replace("[!TIP]", "")
                
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(7.0)
            set_callout_borders(tbl, border_color)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r_icon = p.add_run(icon_title + "\n")
            r_icon.font.bold = True
            r_icon.font.size = Pt(9.5)
            r_icon.font.color.rgb = RGBColor.from_string(border_color)
            
            for al_part in full_alert.strip().split("\n"):
                if not al_part.strip():
                    continue
                p2 = cell.add_paragraph()
                p2.paragraph_format.space_after = Pt(2)
                p2.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p2, al_part.strip(), base_font_size=9.5)
                
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue
            
        # Code Blocks (```python, ```json, ```mermaid, or ASCII tables)
        if line.startswith("```"):
            lang = line[3:].strip().lower()
            code_lines = []
            i += 1
            while i < total_lines and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < total_lines:
                i += 1  # consume closing ```
                
            code_content = "\n".join(code_lines)
            trimmed_lines = [l.strip() for l in code_lines if l.strip()]
            
            # Check 1: Is this an ASCII Box Table?
            if trimmed_lines and trimmed_lines[0].startswith("+") and trimmed_lines[0].endswith("+") and any(l.startswith("|") for l in trimmed_lines):
                # Parse ASCII Box Table
                raw_rows = []
                for l in trimmed_lines:
                    if l.startswith("+") and l.endswith("+"):
                        continue
                    if l.startswith("|") and l.endswith("|"):
                        cells = [c.strip() for c in l.split("|")[1:-1]]
                        raw_rows.append(cells)
                
                if raw_rows:
                    tbl_title = None
                    if len(raw_rows[0]) == 1 or (len(raw_rows[0]) > 1 and all(c == "" for c in raw_rows[0][1:])):
                        tbl_title = raw_rows[0][0]
                        raw_rows = raw_rows[1:]
                        
                    if raw_rows:
                        # If a title was extracted, add as bold caption
                        if tbl_title:
                            p_cap = doc.add_paragraph()
                            p_cap.paragraph_format.space_before = Pt(8)
                            p_cap.paragraph_format.space_after = Pt(3)
                            p_cap.paragraph_format.keep_with_next = True
                            r_cap = p_cap.add_run(f"TABLE: {tbl_title.upper()}")
                            r_cap.font.bold = True
                            r_cap.font.size = Pt(10)
                            r_cap.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)

                        header_row = raw_rows[0]
                        data_rows = raw_rows[1:]
                        num_cols = len(header_row)
                        num_rows = len(data_rows) + 1
                        
                        tbl = doc.add_table(rows=num_rows, cols=num_cols)
                        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                        set_table_borders(tbl, border_color="CBD5E1", header_color="0F2942")
                        
                        col_w = Inches(7.0 / num_cols)
                        
                        # Populate Header
                        set_repeat_header(tbl.rows[0])
                        prevent_row_split(tbl.rows[0])
                        for c_idx, cell in enumerate(tbl.rows[0].cells):
                            cell.width = col_w
                            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                            set_cell_background(cell, "0F2942")  # Deep Navy Header
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_after = Pt(0)
                            r = p.add_run(header_row[c_idx] if c_idx < len(header_row) else "")
                            r.font.bold = True
                            r.font.size = Pt(9)
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            
                        # Populate Data Rows
                        for r_idx, d_row in enumerate(data_rows):
                            row = tbl.rows[r_idx + 1]
                            prevent_row_split(row)
                            bg_hex = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
                            for c_idx, cell in enumerate(row.cells):
                                cell.width = col_w
                                set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                                set_cell_background(cell, bg_hex)
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_after = Pt(0)
                                cell_val = d_row[c_idx] if c_idx < len(d_row) else ""
                                add_formatted_runs(p, cell_val, base_font_size=9)
                                
                        doc.add_paragraph().paragraph_format.space_after = Pt(6)
                        continue

            # Check 2: Is this a Mermaid diagram?
            if lang == "mermaid" or "graph " in code_content or "flowchart " in code_content:
                # Render as structured workflow box
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.rows[0].cells[0]
                cell.width = Inches(7.0)
                set_callout_borders(tbl, "2563EB")
                set_cell_background(cell, "F0F9FF")  # Sky 50
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                p = cell.paragraphs[0]
                r_title = p.add_run("📊 DIAGRAM WORKFLOW & ARCHITECTURE SPECIFICATION\n")
                r_title.font.bold = True
                r_title.font.size = Pt(9.5)
                r_title.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)  # Sky 600
                
                # Add cleaned content
                p_code = cell.add_paragraph()
                p_code.paragraph_format.line_spacing = 1.05
                p_code.paragraph_format.space_after = Pt(0)
                r_c = p_code.add_run(code_content.strip())
                r_c.font.name = 'Consolas'
                r_c.font.size = Pt(8.5)
                r_c.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                continue

            # Check 3: Standard Code Block (Python, JSON, Plain Text)
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(7.0)
            set_code_borders(tbl, "CBD5E1")
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            if lang:
                r_lang = p.add_run(f"// {lang.upper()} CODE SNIPPET\n")
                r_lang.font.name = 'Consolas'
                r_lang.font.bold = True
                r_lang.font.size = Pt(8)
                r_lang.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                
            r = p.add_run(code_content)
            r.font.name = 'Consolas'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Markdown Pipe Tables (| col1 | col2 |)
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Filter out separator row (| :--- | :--- |)
            cleaned_rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue
                cleaned_rows.append(cells)
                
            if cleaned_rows:
                num_cols = len(cleaned_rows[0])
                num_rows = len(cleaned_rows)
                
                tbl = doc.add_table(rows=num_rows, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, border_color="CBD5E1", header_color="0F2942")
                
                # Column widths totaling 7.0 inches
                col_w = Inches(7.0 / num_cols)
                
                for r_idx, row in enumerate(tbl.rows):
                    prevent_row_split(row)
                    if r_idx == 0:
                        set_repeat_header(row)
                    for c_idx, cell in enumerate(row.cells):
                        cell.width = col_w
                        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        
                        if r_idx == 0:
                            set_cell_background(cell, "0F2942")  # Deep Navy Header
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, "FFFFFF")
                        else:
                            set_cell_background(cell, "F8FAFC")  # Subtle Slate Tint
                            
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.1
                        
                        cell_text = cleaned_rows[r_idx][c_idx] if c_idx < len(cleaned_rows[r_idx]) else ""
                        if r_idx == 0:
                            # Header styling
                            r = p.add_run(clean_latex(cell_text.replace("**", "")))
                            r.font.bold = True
                            r.font.size = Pt(9)
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            add_formatted_runs(p, cell_text, base_font_size=9)
                            
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue
            
        # Display Math ($$...$$)
        if line.startswith("$$") and line.endswith("$$") and len(line) > 2:
            math_expr = line[2:-2].strip()
            clean_math = clean_latex(math_expr)
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            
            r = p.add_run(clean_math)
            r.font.name = 'Cambria Math'
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
            i += 1
            continue

        # Bullet Lists (- or *)
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            r_bullet = p.add_run("• ")
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            
            add_formatted_runs(p, line[2:].strip(), base_font_size=10)
            i += 1
            continue

        # Numbered Lists (1. or 2.)
        match_num = re.match(r'^(\d+)\.\s+(.*)$', line)
        if match_num:
            num_str, rest = match_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            r_num = p.add_run(f"{num_str}. ")
            r_num.font.bold = True
            r_num.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
            
            add_formatted_runs(p, rest.strip(), base_font_size=10)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, line, base_font_size=10)
        i += 1

    print(f"Saving compiled Word document to {DOCX_PATH}...")
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    file_size = DOCX_PATH.stat().st_size
    print(f"Successfully generated {DOCX_PATH} ({file_size:,} bytes).")

if __name__ == "__main__":
    generate_docx()
